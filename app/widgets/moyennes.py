"""Widget Moyennes (T1/T2/T3 + Annuelle)."""
from __future__ import annotations

import qtawesome as qta
from qtpy.QtCore import Qt
from qtpy.QtWidgets import (
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QMessageBox,
    QPushButton,
    QTabWidget,
    QTableWidget,
    QVBoxLayout,
    QWidget,
)

from .common import Panel, ProgressOverlay, SearchBar, TopContainer, install_table_sort, set_kind, show_toast
from .table_helpers import (
    fmt_value as _fmt,
    make_grade_item,
    make_item as _v,
    wrap_layout as _wrap,
)
from ..db import Database
from ..models import compute_annuelle, compute_trimestre


class MoyennesWidget(QWidget):
    """Onglet Moyennes : 3 sous-onglets T1/T2/T3 + 1 Annuelle."""

    def __init__(self, db: Database, parent=None) -> None:
        super().__init__(parent)
        self.db = db
        self._build()
        self.refresh()

    def _build(self) -> None:
        root = QVBoxLayout(self)
        root.setContentsMargins(14, 10, 14, 10)
        root.setSpacing(14)

        # --- En-tête de page ---
        self._top = TopContainer(
            "MY",
            "Moyennes",
            "Trimestrielles et annuelle",
        )
        self.btn_export = QPushButton(qta.icon("fa5s.file-export"), "Exporter tout")
        self.btn_export.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_export.setToolTip("Exporter toutes les moyennes vers Excel")
        set_kind(self.btn_export, "tonal")
        self.btn_export.clicked.connect(self._export_all)
        self._top.add_action(self.btn_export)
        root.addWidget(self._top)

        # --- Panel avec sous-onglets ---
        panel = Panel(icon=qta.icon("fa5s.chart-line"))
        # SearchBar (bulletin premium) — filtre live sur la table active
        self.search = SearchBar("Filtrer (nom, n°)…")
        self.search.textChanged.connect(self._apply_filter)
        panel.body_layout().addWidget(self.search)

        self.tabs = QTabWidget()
        self.tables: dict[str, QTableWidget] = {}
        # --- T1 / T2 / T3 : 1 MJ + 1 Compo par trimestre
        for t in (1, 2, 3):
            tbl = QTableWidget(0, 9)
            tbl.setHorizontalHeaderLabels(
                ["N°", "Prénoms", "M.J. (×1)", "Comp. (×2)", "Total M.J", "Total C", "Total/40", "Moyenne/20", "Rang"]
            )
            tbl.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
            tbl.setEditTriggers(tbl.EditTrigger.NoEditTriggers)
            tbl.setAlternatingRowColors(True)
            self.tabs.addTab(tbl, f"Trimestre {t}")
            self.tables[f"T{t}"] = tbl
            # Tri bulletin premium (toutes les colonnes numériques)
            install_table_sort(
                tbl,
                numeric_columns={0, 2, 3, 4, 5, 6, 7, 8},
            )

        # --- Annuelle : 6 notes individuelles + total + MGA + rang
        tbl_a = QTableWidget(0, 11)
        tbl_a.setHorizontalHeaderLabels(
            ["N°", "Prénoms",
             "M.J.1 (×1)", "Comp.1 (×2)",
             "M.J.2 (×1)", "Comp.2 (×2)",
             "M.J.3 (×1)", "Comp.3 (×2)",
             "TOTAL/120", "MGA/20", "Rang"]
        )
        tbl_a.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        tbl_a.setEditTriggers(tbl_a.EditTrigger.NoEditTriggers)
        tbl_a.setAlternatingRowColors(True)
        self.tabs.addTab(tbl_a, "Annuelle")
        self.tables["Annuelle"] = tbl_a
        install_table_sort(
            tbl_a,
            numeric_columns={0, 2, 3, 4, 5, 6, 7, 8, 9, 10},
        )

        panel.body_layout().addWidget(self.tabs)
        root.addWidget(panel, 1)

    # ------------------------------------------------------------------
    def refresh(self) -> None:
        for t in (1, 2, 3):
            self._fill_trim(t)
        self._fill_annuelle()
        # Réapplique le filtre courant
        if hasattr(self, "search") and self.search.text():
            self._apply_filter(self.search.text())

    def _apply_filter(self, text: str) -> None:
        """Filtre la table active (onglet courant) selon ``text``."""
        text = (text or "").strip().lower()
        tbl = self.tables.get(self.tabs.tabText(self.tabs.currentIndex()))
        if tbl is None:
            return
        for r in range(tbl.rowCount()):
            it = tbl.item(r, 1)  # Prénoms
            num_it = tbl.item(r, 0)  # N°
            name = (it.text() if it else "").lower()
            num = (num_it.text() if num_it else "").lower()
            match = (not text) or (text in name) or (text in num)
            tbl.setRowHidden(r, not match)

    def _fill_trim(self, t: int) -> None:
        tbl = self.tables[f"T{t}"]
        rows = compute_trimestre(self.db, t)
        tbl.setRowCount(len(rows))
        if not rows:
            from .table_helpers import make_item as _empty_item
            tbl.setColumnCount(1)
            tbl.setRowCount(1)
            tbl.setItem(0, 0, _empty_item(
                "Aucune donnée — Saisissez des notes dans les onglets matières",
                center=True,
            ))
            return
        for i, r in enumerate(rows):
            tbl.setItem(i, 0, _v(r.student.num, center=True))
            tbl.setItem(i, 1, _v(r.student.full_name))
            tbl.setItem(i, 2, make_grade_item(r.m_moy))
            tbl.setItem(i, 3, make_grade_item(r.c_moy))
            tbl.setItem(i, 4, _v(_fmt(r.m_total) if r.m_total is not None else "—", center=True))
            tbl.setItem(i, 5, _v(_fmt(r.c_total) if r.c_total is not None else "—", center=True))
            # Total brut = somme des 2 moyennes (sur 40 max)
            tbl.setItem(i, 6, _v(_fmt(r.tot), center=True, bold=True))
            # Moy pondérée = (M×1 + C×2) / 3
            tbl.setItem(i, 7, make_grade_item(r.moy))
            tbl.setItem(i, 8, _v(str(r.rank) if r.rank else "NC", center=True, bold=True))
        tbl.resizeColumnsToContents()

    def _fill_annuelle(self) -> None:
        tbl = self.tables["Annuelle"]
        rows = compute_annuelle(self.db)
        tbl.setRowCount(len(rows))
        if not rows:
            from .table_helpers import make_item as _empty_item
            tbl.setColumnCount(1)
            tbl.setRowCount(1)
            tbl.setItem(0, 0, _empty_item(
                "Aucune donnée — Saisissez des notes dans les onglets matières",
                center=True,
            ))
            return
        for i, r in enumerate(rows):
            tbl.setItem(i, 0, _v(r.student.num, center=True))
            tbl.setItem(i, 1, _v(r.student.full_name))
            # 6 notes individuelles : M1, C1, M2, C2, M3, C3
            for c, val in enumerate([r.mj1, r.co1, r.mj2, r.co2, r.mj3, r.co3], 2):
                tbl.setItem(i, c, make_grade_item(val))
            # TOTAL brut = somme des 6 notes (sur 120 max)
            tbl.setItem(i, 8, _v(_fmt(r.total), center=True, bold=True))
            # MGA = moyenne pondérée ((MJ×1 + C×2) / 9) par élève
            tbl.setItem(i, 9, make_grade_item(r.mga))
            tbl.setItem(i, 10, _v(str(r.rank) if r.rank else "NC", center=True, bold=True))
        tbl.resizeColumnsToContents()

    # ------------------------------------------------------------------
    def _export_all(self) -> None:
        from .export_settings import ensure_export_dir
        out_dir = ensure_export_dir(self.db, "export_moyennes_path")
        folder = QFileDialog.getExistingDirectory(self, "Dossier d'export", str(out_dir))
        if not folder:
            return
        overlay = ProgressOverlay(self, "Export des moyennes…")
        overlay.start()
        try:
            import openpyxl  # noqa: WPS433

            for name, tbl in self.tables.items():
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = name
                for c in range(tbl.columnCount()):
                    ws.cell(row=1, column=c + 1, value=tbl.horizontalHeaderItem(c).text())
                for r in range(tbl.rowCount()):
                    for c in range(tbl.columnCount()):
                        it = tbl.item(r, c)
                        if it and it.text():
                            ws.cell(row=r + 2, column=c + 1, value=it.text())
                wb.save(f"{folder}/Moyenne_{name}.xlsx")
            show_toast(self, f"Exports terminés dans {folder}", "success")
        except Exception as e:  # noqa: BLE001
            show_toast(self, "Erreur lors de l'export des moyennes. Vérifiez que les fichiers ne sont pas ouverts dans un autre programme.", "error")
        finally:
            overlay.stop()
